"""Отчёты: экран (JSON), Excel и PDF по одной структуре, цифры совпадают."""

from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Order
from app.models.enums import OrderStatus, Role
from app.services.reports import builders
from app.services.reports.model import window
from app.services.reports.pdf import render_html
from tests.conftest import EmployeeFactory, auth_header
from tests.test_seed import _seed


@pytest.fixture
async def boss(session: AsyncSession, make_employee: EmployeeFactory):  # type: ignore[no-untyped-def]
    await _seed(session)
    return await make_employee(Role.BOSS, login="boss_r")


@pytest.mark.parametrize("kind", ["orders", "rating", "materials", "downtime"])
async def test_each_report_as_json_and_excel(client: AsyncClient, boss, kind: str) -> None:  # type: ignore[no-untyped-def]
    resp = await client.get(f"/api/reports/{kind}?period=month", headers=auth_header(boss))
    assert resp.status_code == 200, resp.text
    report = resp.json()
    assert report["tables"] and report["tables"][0]["rows"], report["title"]
    assert report["period_label"].startswith("30 дней")

    resp = await client.get(
        f"/api/reports/{kind}?period=month&format=xlsx", headers=auth_header(boss)
    )
    assert resp.status_code == 200
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]
    book = load_workbook(BytesIO(resp.content))
    assert book.sheetnames[0] == "Сводка"
    assert book["Сводка"]["A1"].value == report["title"]
    first = book[book.sheetnames[1]]
    table = report["tables"][0]
    note_rows = 2 if table["note"] else 0  # пустая строка и пояснение под таблицей
    assert first.max_row == len(table["rows"]) + 1 + note_rows  # шапка + строки


async def test_orders_report_summary_and_kpis(session: AsyncSession, boss) -> None:  # type: ignore[no-untyped-def]
    report = await builders.orders_report(session, window("month"), None)
    labels = [k.label for k in report.kpis]
    assert labels[:3] == ["Выдано", "Исполнено", "Просрочено"]
    assert report.summary and report.summary.startswith("Выдано нарядов:")
    assert report.summary_source == "rules"


async def test_materials_report_shows_overuse_with_norm(session: AsyncSession, boss) -> None:  # type: ignore[no-untyped-def]
    report = await builders.materials_report(session, window("month"), None)
    over = report.tables[0].rows
    assert over, "в сиде заложен перерасход — он должен попасть в отчёт"
    top = over[0]
    assert top["ratio"].startswith("×") and top["fact"] and top["norm"]


async def test_pdf_template_renders_cyrillic_and_logo(session: AsyncSession, boss) -> None:  # type: ignore[no-untyped-def]
    report = await builders.downtime_report(session, window("month"), None)
    html = render_html(report)
    assert "Простои оборудования" in html and "km-logo-white.png" in html
    assert "counter(pages)" in html


async def test_pdf_when_pango_available(session: AsyncSession, boss) -> None:  # type: ignore[no-untyped-def]
    try:
        import weasyprint  # noqa: F401
    except OSError:
        pytest.skip("нет системных библиотек Pango (Windows без GTK) — PDF проверяется в CI")
    from app.services.reports.pdf import to_pdf

    data = to_pdf(await builders.rating_report(session, window("month")))
    assert data.startswith(b"%PDF")


async def test_order_report_visible_to_its_worker_only(
    client: AsyncClient,
    session: AsyncSession,
    boss,
    make_employee: EmployeeFactory,  # type: ignore[no-untyped-def]
) -> None:
    order = await session.scalar(
        select(Order).where(Order.status == OrderStatus.CLOSED).order_by(Order.id).limit(1)
    )
    assert order is not None
    resp = await client.get(f"/api/reports/order/{order.id}?format=json", headers=auth_header(boss))
    assert resp.status_code == 200
    body = resp.json()
    assert body["title"] == f"Наряд №{order.number}"
    assert [t["title"] for t in body["tables"]][:2] == ["Наряд", "Ход работ"]

    stranger = await make_employee(Role.WORKER, login="stranger")
    resp = await client.get(
        f"/api/reports/order/{order.id}?format=json", headers=auth_header(stranger)
    )
    assert resp.status_code == 403


async def test_custom_period_validation(client: AsyncClient, boss) -> None:  # type: ignore[no-untyped-def]
    resp = await client.get(
        "/api/reports/orders?period=custom&date_from=2026-10-05&date_to=2026-10-01",
        headers=auth_header(boss),
    )
    assert resp.status_code == 422
    assert "поменяйте" in resp.json()["detail"].lower()
