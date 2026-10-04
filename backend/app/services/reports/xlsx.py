"""Отчёт в Excel: каждая таблица — свой лист, первый лист — сводка с показателями."""

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.reports.model import Report, Table, fmt_dt

# Палитра дизайн-системы (docs/DESIGN.md)
BAR = "0B3F8C"
INK = "172230"
INK_3 = "657080"
PLATE = "EDF0F4"
LINE = "DFE3E9"
RED = "C8281C"
RED_SOFT = "FDEDEB"
GREEN = "17723A"

THIN = Side(style="thin", color=LINE)
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet_name(title: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Лист"
    base, n = name, 2
    while name in used:
        name = f"{base[:28]} {n}"
        n += 1
    used.add(name)
    return name


def _write_table(ws: Worksheet, table: Table, start_row: int) -> int:
    for col, c in enumerate(table.columns, 1):
        cell = ws.cell(row=start_row, column=col, value=c.title)
        cell.font = Font(bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=PLATE)
        cell.border = GRID
        cell.alignment = Alignment(
            horizontal="right" if c.align == "right" else "left", vertical="center", wrap_text=True
        )
        ws.column_dimensions[get_column_letter(col)].width = c.width
    for r, row in enumerate(table.rows, start_row + 1):
        flagged = bool(table.highlight and row.get(table.highlight))
        for col, c in enumerate(table.columns, 1):
            cell = ws.cell(row=r, column=col, value=row.get(c.key, ""))
            cell.border = GRID
            cell.alignment = Alignment(
                horizontal="right" if c.align == "right" else "left",
                vertical="top",
                wrap_text=c.width >= 30,
            )
            if flagged:
                cell.fill = PatternFill("solid", fgColor=RED_SOFT)
    end = start_row + len(table.rows)
    if table.note:
        end += 2
        ws.cell(row=end, column=1, value=table.note).font = Font(italic=True, color=INK_3)
    return end


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    used: set[str] = set()
    ws.title = _sheet_name("Сводка", used)

    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BAR)
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 28
    ws["A2"] = report.period_label + (f" · {report.filters}" if report.filters else "")
    ws["A2"].font = Font(color=INK_3)
    ws["A3"] = f"АО «Костанайские Минералы» · НарядAI · сформирован {fmt_dt(report.generated_at)}"
    ws["A3"].font = Font(color=INK_3, size=9)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22

    row = 5
    for kpi in report.kpis:
        ws.cell(row=row, column=1, value=kpi.label).font = Font(color=INK_3)
        value = ws.cell(row=row, column=2, value=kpi.value)
        color = {"danger": RED, "ok": GREEN}.get(kpi.tone, INK)
        value.font = Font(bold=True, size=13, color=color)
        row += 1
    if report.summary:
        row += 1
        ws.cell(row=row, column=1, value="Сводка").font = Font(bold=True, color=INK)
        cell = ws.cell(row=row + 1, column=1, value=report.summary)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
        ws.row_dimensions[row + 1].height = max(45, 15 * (len(report.summary) // 90 + 1))

    for table in report.tables:
        sheet = wb.create_sheet(_sheet_name(table.title, used))
        _write_table(sheet, table, 1)
        sheet.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
